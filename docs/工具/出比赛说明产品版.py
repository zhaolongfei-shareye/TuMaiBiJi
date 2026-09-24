#!/usr/bin/env python3
"""把 docs/比赛说明-产品版.md 一次出成 .docx（可编辑提交件）与 .pdf（排版件）。

两条渲染共用同一份 markdown，避免正文漂移：
  · docx 走 python-docx，沿用 docs/工具/出比赛说明docx.py 的解析口径；
  · pdf 走「markdown → 带打印样式的 HTML → headless Chrome 打印」，图注与分栏靠 CSS 控。
只覆盖这份文档实际用到的语法：# 三级标题、**加粗**、> 引用、- 列表、表格、代码块、--- 分隔线、![](图)。
"""
import html as htmlmod
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SRC = Path('/Users/zlfmac/Documents/TuMaiBiJi/docs/比赛说明-产品版.md')
OUT_DOCX = Path('/Users/zlfmac/Documents/TuMaiBiJi/比赛/图麦笔记-产品说明-20260924.docx')
OUT_PDF = Path('/Users/zlfmac/Documents/TuMaiBiJi/比赛/图麦笔记-产品说明-20260924.pdf')
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'

INK = RGBColor(0x23, 0x25, 0x2C)
SUB = RGBColor(0x5C, 0x60, 0x68)
ACCENT = RGBColor(0x3F, 0x52, 0xD6)

lines = SRC.read_text(encoding='utf-8').split('\n')


def blocks(ls):
    """把 markdown 切成 (kind, payload) 序列，两种渲染各自消费。"""
    out, i = [], 0
    while i < len(ls):
        ln = ls[i].rstrip()
        if not ln.strip() or ln.strip() == '>':
            i += 1
            continue
        if ln.strip().startswith('```'):
            i += 1
            buf = []
            while i < len(ls) and not ls[i].strip().startswith('```'):
                buf.append(ls[i])
                i += 1
            i += 1
            out.append(('code', buf))
            continue
        if ln.strip() == '---':
            out.append(('rule', None))
            i += 1
            continue
        m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)$', ln)
        if m:
            out.append(('img', (SRC.parent / m.group(1)).resolve()))
            i += 1
            continue
        m = re.match(r'^\*([^*].*?)\*$', ln)
        if m:
            out.append(('caption', m.group(1)))
            i += 1
            continue
        m = re.match(r'^(#{1,3})\s+(.*)$', ln)
        if m:
            out.append(('h' + str(len(m.group(1))), m.group(2)))
            i += 1
            continue
        if ln.startswith('|'):
            rows = []
            while i < len(ls) and ls[i].strip().startswith('|'):
                cells = [c.strip() for c in ls[i].strip().strip('|').split('|')]
                if not re.fullmatch(r'[-: ]+', cells[0] or '-'):
                    rows.append(cells)
                i += 1
            out.append(('table', rows))
            continue
        if ln.startswith('> '):
            buf = []
            while i < len(ls) and (ls[i].startswith('> ') or ls[i].strip() == '>'):
                if ls[i].startswith('> '):
                    buf.append(ls[i][2:])
                i += 1
            out.append(('quote', ' '.join(buf)))
            continue
        m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
        if m:
            out.append(('li', m.group(2)))
            i += 1
            continue
        out.append(('p', ln))
        i += 1
    return out


def inline(text):
    """拆成 [(文本, 加粗)]；行内反引号只去掉符号——Word 与 PDF 里等宽字体反而突兀。"""
    text = text.replace('`', '')
    return [(c, k % 2 == 1) for k, c in enumerate(re.split(r'\*\*(.+?)\*\*', text)) if c]


# ---------------------------------------------------------------- docx
doc = Document()
st = doc.styles['Normal']
st.font.name = 'PingFang SC'
st.font.size = Pt(10.5)
st.font.color.rgb = INK

for kind, payload in blocks(lines):
    if kind == 'rule':
        continue
    if kind.startswith('h'):
        h = doc.add_heading(level={'h1': 0, 'h2': 1, 'h3': 2}[kind])
        for text, bold in inline(payload):
            r = h.add_run(text)
            r.bold = bold
            r.font.color.rgb = ACCENT if kind in ('h1', 'h2') else INK
        continue
    if kind == 'code':
        for q in payload:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(q.replace('`', ''))
            r.font.name = 'Menlo'
            r.font.size = Pt(8.5)
            r.font.color.rgb = SUB
        continue
    if kind == 'img':
        if not payload.exists():
            raise SystemExit(f'配图文件不存在：{payload}')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(payload), width=Inches(6.2))
        continue
    if kind == 'caption':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(payload)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = SUB
        continue
    if kind == 'table':
        t = doc.add_table(rows=len(payload), cols=max(len(r) for r in payload))
        t.style = 'Light Grid Accent 1'
        for ri, cells in enumerate(payload):
            for ci, val in enumerate(cells):
                cell = t.cell(ri, ci)
                cell.text = ''
                p = cell.paragraphs[0]
                for text, bold in inline(val):
                    r = p.add_run(text)
                    r.bold = bold or ri == 0
                    r.font.size = Pt(9)
        doc.add_paragraph()
        continue
    if kind == 'quote':
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        for text, bold in inline(payload):
            r = p.add_run(text)
            r.italic = True
            r.bold = bold
            r.font.color.rgb = SUB
        continue
    if kind == 'li':
        p = doc.add_paragraph(style='List Bullet')
        for text, bold in inline(payload):
            r = p.add_run(text)
            r.bold = bold
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for text, bold in inline(payload):
        r = p.add_run(text)
        r.bold = bold

OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT_DOCX)
print(OUT_DOCX, OUT_DOCX.stat().st_size, 'bytes')

# ---------------------------------------------------------------- pdf
CSS = """
@page { size: A4; margin: 17mm 16mm 16mm; }
body { font: 10.5pt/1.85 "Songti SC","PingFang SC",serif; color:#23252C; margin:0 }
h1 { font: 600 21pt/1.4 "PingFang SC"; color:#3F52D6; margin:0 0 4pt; page-break-before:avoid }
h1:first-of-type { font-size: 24pt }
h2 { font: 600 14pt/1.5 "PingFang SC"; color:#3F52D6; margin:22pt 0 7pt;
     border-bottom:.6pt solid #D8D5CB; padding-bottom:4pt; page-break-after:avoid }
h3 { font: 600 11.5pt/1.5 "PingFang SC"; margin:14pt 0 5pt; page-break-after:avoid }
p { margin:0 0 8pt; text-align:left; overflow-wrap:anywhere }
b { font-weight:600 }
blockquote { margin:0 0 10pt; padding:8pt 12pt; background:#F4F2EC; border-left:2.5pt solid #3F52D6;
             color:#4B4E55; font-style:italic }
ul { margin:0 0 9pt; padding-left:16pt } li { margin-bottom:4pt; text-align:left }
table { width:100%; border-collapse:collapse; margin:6pt 0 12pt; font:8.6pt/1.6 "PingFang SC";
        table-layout:fixed }
th { background:#EEF1FA; text-align:left } th,td { border:.5pt solid #C9CCD4; padding:4.5pt 6pt;
        vertical-align:top; word-break:break-word }
code { font-family:Menlo,monospace; font-size:8.4pt; word-break:break-all }
pre { background:#F4F2EC; border:.5pt solid #DDD9CF; padding:8pt 10pt; font:8.4pt/1.7 Menlo,monospace;
      white-space:pre-wrap; margin:6pt 0 12pt }
figure { margin:10pt 0 14pt; page-break-inside:avoid }
figure img { width:100%; display:block; border:.5pt solid #DDD9CF }
figcaption { font:8.4pt/1.6 "PingFang SC"; color:#5C6068; font-style:italic; margin-top:5pt;
             text-align:center }
"""


def to_html(bs):
    out = []
    for kind, payload in bs:
        if kind == 'rule':
            continue
        if kind.startswith('h'):
            tag = kind
            txt = ''.join(f'<b>{htmlmod.escape(t)}</b>' if b else htmlmod.escape(t)
                          for t, b in inline(payload))
            out.append(f'<{tag}>{txt}</{tag}>')
        elif kind == 'code':
            out.append('<pre>' + htmlmod.escape('\n'.join(payload)) + '</pre>')
        elif kind == 'img':
            out.append(f'<figure><img src="{payload.as_uri()}"></figure>')
        elif kind == 'caption':
            out.append(f'<figcaption>{htmlmod.escape(payload)}</figcaption>')
        elif kind == 'table':
            rows = ['<table>']
            for ri, cells in enumerate(payload):
                tag = 'th' if ri == 0 else 'td'
                cells_html = ''.join(
                    '<%s>%s</%s>' % (tag,
                                     ''.join(f'<b>{htmlmod.escape(t)}</b>' if b else htmlmod.escape(t)
                                             for t, b in inline(c)), tag)
                    for c in cells)
                rows.append('<tr>' + cells_html + '</tr>')
            rows.append('</table>')
            out.append(''.join(rows))
        elif kind == 'quote':
            out.append('<blockquote>' + ''.join(
                f'<b>{htmlmod.escape(t)}</b>' if b else htmlmod.escape(t)
                for t, b in inline(payload)) + '</blockquote>')
        elif kind == 'li':
            out.append('<ul><li>' + ''.join(
                f'<b>{htmlmod.escape(t)}</b>' if b else htmlmod.escape(t)
                for t, b in inline(payload)) + '</li></ul>')
        else:
            out.append('<p>' + ''.join(
                f'<b>{htmlmod.escape(t)}</b>' if b else htmlmod.escape(t)
                for t, b in inline(payload)) + '</p>')
    # 图注紧跟图：把 figcaption 塞进上一个 figure；相邻 <ul> 合成一个列表
    merged = []
    for frag in out:
        if merged and frag.startswith('<figcaption>') and merged[-1].endswith('</figure>'):
            merged[-1] = merged[-1][: -len('</figure>')] + frag + '</figure>'
        elif merged and frag.startswith('<ul>') and merged[-1].endswith('</ul>'):
            merged[-1] = merged[-1][: -len('</ul>')] + frag[len('<ul>'):]
        else:
            merged.append(frag)
    return ('<!doctype html><html lang=zh><head><meta charset=utf-8>'
            '<title>图麦笔记 · 产品说明</title><style>' + CSS + '</style></head><body>'
            + '\n'.join(merged) + '</body></html>')


bs = blocks(lines)
tmp_html = Path('/tmp/docwork/产品版.html')
tmp_html.parent.mkdir(exist_ok=True)
tmp_html.write_text(to_html(bs), encoding='utf-8')
subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--no-first-run',
                '--no-pdf-header-footer', f'--print-to-pdf={OUT_PDF}',
                tmp_html.as_uri()], check=True, capture_output=True)
print(OUT_PDF, OUT_PDF.stat().st_size, 'bytes')
